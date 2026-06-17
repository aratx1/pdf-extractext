"""PDF Summary API router with dependency injection."""

import io
import re
from uuid import UUID
import httpx
from docx import Document
from docx.shared import Pt, RGBColor
from fastapi import APIRouter, File, UploadFile, HTTPException, Depends
from fastapi.responses import StreamingResponse
from app.application.services.summary_service import SummaryService
from app.infrastructure.external.openrouter_client import MissingAPIKeyError, OpenRouterError
from app.presentation.schemas.pdf_summary import (
    SummaryResponse,
    SummaryListResponse,
    HealthResponse,
)

router = APIRouter(prefix="/api", tags=["summaries"])


def get_summary_service() -> SummaryService:
    from app.main import get_summary_service as _get_service

    return _get_service()


@router.post("/summarize", response_model=SummaryResponse)
async def summarize_pdf(
    file: UploadFile = File(...),
    service: SummaryService = Depends(get_summary_service),
):
    if not file.filename.endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Solo podemos leer archivos PDF.")

    content = await file.read()
    if len(content) == 0:
        raise HTTPException(status_code=400, detail="Adjuntaste un archivo vacio.")

    try:
        summary = await service.create_summary(content, file.filename)
    except MissingAPIKeyError:
        raise HTTPException(
            status_code=502,
            detail=(
                "Se necesita generar una API key en el .env. "
                "Crea una gratis en https://openrouter.ai/keys y agrégala como "
                "OPENROUTER_API_KEY=... (y opcionalmente OPENROUTER_MODEL con un model ID válido)."
            ),
        )
    except OpenRouterError as exc:
        raise HTTPException(
            status_code=502,
            detail=str(exc),
        )
    except httpx.TimeoutException:
        raise HTTPException(
            status_code=504,
            detail="El modelo de IA tardó demasiado en responder. Intenta con un PDF más corto.",
        )
    except httpx.HTTPStatusError as exc:
        if exc.response.status_code in (400, 401, 402):
            raise HTTPException(
                status_code=502,
                detail=(
                    "Se necesita generar una API key en el .env. "
                    "Crea una gratis en https://openrouter.ai/keys y agrégala como "
                    "OPENROUTER_API_KEY=... (y opcionalmente OPENROUTER_MODEL con un model ID válido)."
                ),
            )
        raise HTTPException(
            status_code=503,
            detail=f"Error del proveedor de IA ({exc.response.status_code}): {exc}",
        )
    except httpx.HTTPError as exc:
        raise HTTPException(
            status_code=503,
            detail=f"No se pudo contactar al proveedor de IA: {exc}",
        )
    return SummaryResponse(
        id=summary.id,
        original_filename=summary.original_filename,
        summary_text=summary.summary_text,
        created_at=summary.created_at,
    )


@router.get("/summaries", response_model=SummaryListResponse)
async def list_summaries(
    limit: int = 100,
    service: SummaryService = Depends(get_summary_service),
):
    summaries = await service.list_summaries(limit)
    return SummaryListResponse(
        summaries=[
            SummaryResponse(
                id=s.id,
                original_filename=s.original_filename,
                summary_text=s.summary_text,
                created_at=s.created_at,
            )
            for s in summaries
        ],
        total=len(summaries),
    )


@router.get("/summaries/{summary_id}", response_model=SummaryResponse)
async def get_summary(
    summary_id: UUID,
    service: SummaryService = Depends(get_summary_service),
):
    summary = await service.get_summary(summary_id)
    if not summary:
        raise HTTPException(status_code=404, detail="Summary not found")
    return SummaryResponse(
        id=summary.id,
        original_filename=summary.original_filename,
        summary_text=summary.summary_text,
        created_at=summary.created_at,
    )


@router.get("/summaries/{summary_id}/export")
async def export_summary_docx(
    summary_id: UUID,
    service: SummaryService = Depends(get_summary_service),
):
    """Export a summary as a .docx file using python-docx."""
    summary = await service.get_summary(summary_id)
    if not summary:
        raise HTTPException(status_code=404, detail="Summary not found")

    doc = Document()

    # Title
    title = doc.add_heading(level=1)
    title_run = title.runs[0] if title.runs else title.add_run("")
    title.clear()
    title_run = title.add_run(f"Resumen: {summary.original_filename}")
    title_run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)  # blue

    doc.add_paragraph()  # spacer

    # Parse the markdown-like text into paragraphs
    lines = summary.summary_text.splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph()
            continue

        # Heading detection: ## Heading or # Heading
        h2_match = re.match(r'^#{1,2}\s+(.*)', stripped)
        if h2_match:
            heading_text = h2_match.group(1).strip()
            h = doc.add_heading(level=2)
            h.clear()
            run = h.add_run(heading_text)
            run.font.color.rgb = RGBColor(0x37, 0x51, 0x6F)
            continue

        # Bullet points: - item or * item
        bullet_match = re.match(r'^[-*]\s+(.*)', stripped)
        if bullet_match:
            para = doc.add_paragraph(style='List Bullet')
            para.add_run(bullet_match.group(1))
            continue

        # Bold text inline: **text**
        para = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', stripped)
        for part in parts:
            bold_match = re.match(r'^\*\*(.*?)\*\*$', part)
            if bold_match:
                run = para.add_run(bold_match.group(1))
                run.bold = True
            else:
                para.add_run(part)

    # Save to buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    safe_name = re.sub(r'[^\w\-.]', '_', summary.original_filename.replace('.pdf', ''))
    filename = f"resumen_{safe_name}.docx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/health", response_model=HealthResponse)
async def health_check(service: SummaryService = Depends(get_summary_service)):
    ai_available = await service._ai_provider.health_check()
    return HealthResponse(
        status="healthy" if ai_available else "degraded",
        ai_provider_available=ai_available,
    )
